"""Build the editable Measurement highlights file."""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


HIGHLIGHTS = [
    "Fidelity scores are used to choose fusion rules without downstream validation.",
    "Test-independent splits retrain detectors with all 300 test images held out.",
    "All 42 clean visible-detector comparisons favor the unfused visible image.",
    "Deficits span 0.046 to 0.246 mAP50; both one-run evaluators survive BH correction.",
    "Evaluator modality reverses gains, so rankings require measurement context.",
]


def set_font(run, name: str, size: float, bold: bool | None = None) -> None:
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor(0, 0, 0)
    if bold is not None:
        run.bold = bold


def main() -> None:
    output = Path(
        r"D:\Code\_server_handoff_v2\ssgf\paper\HIGHLIGHTS_MEASUREMENT.docx"
    )
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.9)
    section.bottom_margin = Inches(0.9)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)

    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing = 1.12

    title_style = document.styles["Title"]
    style_p_pr = title_style._element.get_or_add_pPr()
    style_border = style_p_pr.find(qn("w:pBdr"))
    if style_border is not None:
        style_p_pr.remove(style_border)

    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title.paragraph_format.space_after = Pt(8)
    title_border = title._p.get_or_add_pPr().find(qn("w:pBdr"))
    if title_border is not None:
        title._p.get_or_add_pPr().remove(title_border)
    set_font(title.add_run("Highlights"), "Calibri", 22, True)

    article = document.add_paragraph()
    article.paragraph_format.space_after = Pt(14)
    set_font(
        article.add_run(
            "Evaluator-Dependent Validity of Fidelity Metrics for "
            "Visible-Infrared Image Fusion in Pedestrian Detection"
        ),
        "Calibri",
        12,
        False,
    )
    article.runs[0].italic = True

    introduction = document.add_paragraph()
    introduction.paragraph_format.space_after = Pt(12)
    set_font(
        introduction.add_run(
            "These five statements summarize the measurement contribution and "
            "the principal independent-test findings."
        ),
        "Calibri",
        10.5,
    )

    for text in HIGHLIGHTS:
        if len(text) > 85:
            raise ValueError(f"highlight exceeds 85 characters: {len(text)} {text}")
        paragraph = document.add_paragraph(style="List Bullet")
        paragraph.paragraph_format.left_indent = Inches(0.26)
        paragraph.paragraph_format.first_line_indent = Inches(-0.18)
        paragraph.paragraph_format.space_after = Pt(9)
        set_font(paragraph.add_run(text), "Calibri", 11)

    core = document.core_properties
    core.title = "Highlights"
    core.subject = "Measurement manuscript highlights"
    core.author = "Jianyong Zhao, Zigao Wu, Mengfei Ma"
    core.keywords = "measurement validity, visible-infrared fusion, evaluation protocol"

    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)
    print(output)


if __name__ == "__main__":
    main()
